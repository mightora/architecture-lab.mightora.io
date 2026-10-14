from zipfile import ZipFile
from lxml import etree
from pptx import Presentation
from docx import Document
from pathlib import Path
for path in Path('checks').glob('editable-export.*'):
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith(('.xml', '.rels')):
                etree.fromstring(archive.read(name))
    print(path.name, 'XML validated')
slides = Presentation('checks/editable-export.pptx')
assert len(slides.slides) == 1
assert len(slides.slides[0].shapes) == 6
assert any(s.has_text_frame and s.text == 'Power Apps' for s in slides.slides[0].shapes)
document = Document('checks/editable-export.docx')
assert len(document.inline_shapes) == 1
assert len(document.element.xpath('//w:drawing')) == 1
assert len(document.element.xpath('//w:t[contains(text(), "Power Apps")]')) == 1
print('Native objects and editable labels validated in both Office packages.')
converted = Presentation('checks/svg-converted-export.pptx')
assert len(converted.slides[0].shapes) == 6
assert any(s.has_text_frame and s.text == 'Plan the menu' for s in converted.slides[0].shapes)
converted_word = Document('checks/svg-converted-export.docx')
assert len(converted_word.inline_shapes) == 1
assert converted_word.element.xpath('//w:drawing//w:t[contains(text(), "Plan the menu")]')
print('SVG-derived Word and PowerPoint documents open with native text and drawing objects.')
