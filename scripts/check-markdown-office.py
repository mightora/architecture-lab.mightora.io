from zipfile import ZipFile
from lxml import etree
from docx import Document
path='checks/markdown-document.docx'
with ZipFile(path) as archive:
    for name in archive.namelist():
        if name.endswith(('.xml','.rels')):
            etree.fromstring(archive.read(name))
    assert not any(n.startswith('word/media/') for n in archive.namelist())
doc=Document(path)
assert len(doc.inline_shapes)==9
assert len(doc.tables)==1
assert doc.tables[0].cell(1,0).text=='Host'
assert doc.paragraphs[0].style.style_id=='Heading1'
text='\n'.join(p.text for p in doc.paragraphs)
assert 'Invalid syntax' in text and 'invalid diagram' in text
assert 'ordinary code' in text and 'inline code' in text
assert doc.element.xpath('//w:rPr/w:b')
assert doc.element.xpath('//w:rPr/w:i')
assert len(doc.element.xpath('//w:drawing//w:t[contains(text(), "Editable diagram label")]'))==9
assert doc.element.xpath('//w:numPr')
print('Mixed Markdown DOCX opens with nine native editable diagram groups, heading styles, table cells, lists and retained failure source.')
